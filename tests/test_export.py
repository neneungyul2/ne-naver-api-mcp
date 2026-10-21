import os
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ne_naver_api_mcp.export import (
    ExportResult,
    Sheet,
    export_sheets,
    neutralize_formula,
    sanitize_base_name,
)


def inside(directory: Path, path: str) -> bool:
    return Path(path).resolve().is_relative_to(directory.resolve())


def test_sanitize_base_name_removes_path_characters_but_keeps_korean():
    s = sanitize_base_name("../../evil")
    assert "/" not in s and "\\" not in s and ".." not in s
    assert sanitize_base_name("그래머 비교 2025") == "그래머 비교 2025"
    assert sanitize_base_name("...") == "export"
    assert sanitize_base_name("") == "export"
    assert sanitize_base_name("a:b*c?d") == "a_b_c_d"


def test_sanitize_base_name_replaces_backslashes():
    assert sanitize_base_name("a\\b") == "a_b"
    s = sanitize_base_name("..\\..\\evil")
    assert "\\" not in s and "/" not in s and ".." not in s


def test_neutralize_formula_only_touches_strings_that_look_like_formulas():
    assert neutralize_formula('=HYPERLINK("x")') == "'=HYPERLINK(\"x\")"
    assert neutralize_formula("+1") == "'+1"
    assert neutralize_formula("-1") == "'-1"
    assert neutralize_formula("@SUM") == "'@SUM"
    assert neutralize_formula("\t탭") == "'\t탭"
    assert neutralize_formula("\r줄") == "'\r줄"
    assert neutralize_formula(-3) == -3
    assert neutralize_formula(None) is None
    assert neutralize_formula("정상") == "정상"


def test_export_writes_xlsx_and_one_csv_per_sheet_inside_the_directory(tmp_path):
    res = export_sheets(
        [
            Sheet("트렌드", ["기간", "A"], [["2025-01-01", 12.5], ["2025-02-01", None]]),
            Sheet("한계", ["문구"], [["상대 지수입니다"]]),
        ],
        tmp_path,
        "비교",
    )
    assert isinstance(res, ExportResult)
    assert Path(res.xlsx_path).exists() and len(res.csv_paths) == 2
    for p in [res.xlsx_path, *res.csv_paths]:
        assert Path(p).exists() and inside(tmp_path, p)
    assert res.to_dict() == {"xlsxPath": res.xlsx_path, "csvPaths": res.csv_paths}


def test_base_name_with_path_parts_never_escapes_the_directory(tmp_path):
    target = tmp_path / "out"
    res = export_sheets([Sheet("s", ["a"], [["x"]])], target, "../x")
    for p in [res.xlsx_path, *res.csv_paths]:
        assert inside(target, p)
    res2 = export_sheets([Sheet("s", ["a"], [["x"]])], target, "..\\x")
    for p in [res2.xlsx_path, *res2.csv_paths]:
        assert inside(target, p)


def test_csv_has_bom_escapes_fields_and_keeps_korean(tmp_path):
    res = export_sheets(
        [Sheet("t", ["이름", "메모"], [["그래머", 'a,"b"\nc'], [None, 3]])], tmp_path, "csv"
    )
    raw = Path(res.csv_paths[0]).read_bytes()
    assert raw[:3] == b"\xef\xbb\xbf"
    text = raw.decode("utf-8")[1:]
    assert text.startswith("이름,메모\r\n")
    assert '그래머,"a,""b""\nc"' in text
    assert ",3" in text


def test_csv_cells_are_neutralized_too(tmp_path):
    res = export_sheets([Sheet("t", ["=bad"], [["=1+1"]])], tmp_path, "csv")
    text = Path(res.csv_paths[0]).read_text(encoding="utf-8-sig")
    assert "'=bad" in text and "'=1+1" in text


def test_xlsx_round_trips_and_formula_like_cells_are_neutralized(tmp_path):
    res = export_sheets(
        [Sheet("트렌드", ["기간", "값"], [["2025-01-01", 12.5], ["=1+1", "정상"]])], tmp_path, "x"
    )
    ws = load_workbook(res.xlsx_path)["트렌드"]
    assert ws.cell(1, 1).value == "기간"
    assert ws.cell(2, 2).value == 12.5
    assert ws.cell(3, 1).value == "'=1+1"
    assert ws.cell(3, 1).data_type == "s"  # 수식으로 저장되지 않는다


def test_illegal_control_characters_do_not_break_the_workbook(tmp_path):
    res = export_sheets([Sheet("t", ["a"], [["가\x07나"]])], tmp_path, "ctrl")
    assert load_workbook(res.xlsx_path)["t"].cell(2, 1).value == "가나"


def test_sheet_names_are_cleaned_to_excel_rules(tmp_path):
    long_name = "가" * 40 + "[x]"
    res = export_sheets(
        [Sheet(long_name, ["a"], [["x"]]), Sheet("가\\나", ["a"], [["x"]]), Sheet("가\\나", ["a"], [["y"]])],
        tmp_path,
        "names",
    )
    names = load_workbook(res.xlsx_path).sheetnames
    assert all(len(n) <= 31 for n in names)
    assert not any(c in n for n in names for c in "[]*?:\\/")
    assert len({n.lower() for n in names}) == 3  # 중복 이름도 서로 달라진다


def test_exporting_twice_overwrites_without_error(tmp_path):
    for _ in range(2):
        export_sheets([Sheet("t", ["a"], [["x"]])], tmp_path, "same")
    assert len(list(tmp_path.iterdir())) == 2  # xlsx 1 + csv 1


def test_no_sheets_still_creates_a_valid_workbook(tmp_path):
    res = export_sheets([], tmp_path, "empty")
    assert load_workbook(res.xlsx_path).sheetnames
    assert res.csv_paths == []


def test_files_are_utf8_regardless_of_the_console_locale(tmp_path):
    code = (
        "from ne_naver_api_mcp.export import Sheet, export_sheets\n"
        f"export_sheets([Sheet('한글시트', ['이름'], [['그래머인사이드']])], r'{tmp_path}', '한글')\n"
    )
    env = {**os.environ, "PYTHONUTF8": "0", "PYTHONIOENCODING": "cp949"}
    subprocess.run([sys.executable, "-c", code], env=env, check=True)
    csv_file = next(tmp_path.glob("*.csv"))
    assert "그래머인사이드" in csv_file.read_bytes().decode("utf-8")
    assert load_workbook(next(tmp_path.glob("*.xlsx")))["한글시트"].cell(2, 1).value == "그래머인사이드"
