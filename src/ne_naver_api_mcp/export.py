"""분석 결과를 xlsx 하나와 시트별 CSV(UTF-8 BOM)로 저장한다."""
import csv
import re
from collections.abc import Sequence
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

Cell = str | int | float | None

_FORMULA_START = re.compile(r"[=+\-@\t\r]")


@dataclass
class Sheet:
    name: str
    columns: list[str]
    rows: list[list[Cell]] = field(default_factory=list)


@dataclass
class ExportResult:
    xlsx_path: str
    csv_paths: list[str]

    def to_dict(self) -> dict:
        return {"xlsxPath": self.xlsx_path, "csvPaths": self.csv_paths}


def sanitize_base_name(name: str) -> str:
    """파일명으로 쓸 수 없는 문자(\\ / : * ? " < > |, 제어 문자)와 경로 이동(..)을 `_`로 바꾼다."""
    s = re.sub(r"^[. ]+|[. ]+$", "", name.strip())
    s = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", s)
    s = re.sub(r"\.{2,}", "_", s)
    return "export" if not s or re.fullmatch(r"_+", s) else s


def neutralize_formula(value: Cell) -> Cell:
    """엑셀이 수식으로 읽는 문자열(=, +, -, @, 탭, CR로 시작)은 앞에 '를 붙여 무력화한다."""
    if isinstance(value, str) and _FORMULA_START.match(value):
        return "'" + value
    return value


def _cell(value: Cell) -> Cell:
    value = neutralize_formula(value)
    if isinstance(value, str):
        # openpyxl이 거부하는 제어 문자(웹 글에 섞여 올 수 있다)를 제거한다
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def _sheet_title(name: str) -> str:
    return re.sub(r"[\[\]*?:\\/]", "_", name).strip()[:31] or "Sheet"


def export_sheets(sheets: Sequence[Sheet], directory: Path, base_name: str) -> ExportResult:
    """시트들을 `<base>.xlsx` 하나와 시트별 `<base>-<시트>.csv`로 directory 아래에 저장한다."""
    out_dir = Path(directory).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    base = sanitize_base_name(base_name)

    workbook = Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()
    for sheet in sheets:
        title = _sheet_title(sheet.name)
        n = 2
        while title.lower() in used:
            title = f"{title[:28]}_{n}"
            n += 1
        used.add(title.lower())
        ws = workbook.create_sheet(title)
        ws.append([_cell(c) for c in sheet.columns])
        for row in sheet.rows:
            ws.append([_cell(c) for c in row])
    if not workbook.sheetnames:
        workbook.create_sheet("Sheet")
    xlsx_path = out_dir / f"{base}.xlsx"
    workbook.save(xlsx_path)

    csv_paths: list[str] = []
    for sheet in sheets:
        path = out_dir / f"{base}-{sanitize_base_name(sheet.name)}.csv"
        with path.open("w", encoding="utf-8-sig", newline="") as f:
            writer = csv.writer(f, lineterminator="\r\n")
            writer.writerow([_cell(c) for c in sheet.columns])
            for row in sheet.rows:
                writer.writerow([_cell(c) for c in row])
        csv_paths.append(str(path))
    return ExportResult(str(xlsx_path), csv_paths)
